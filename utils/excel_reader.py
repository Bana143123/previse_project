from openpyxl import load_workbook

def read_excel_data(file_path):
    workbook = load_workbook(filename=file_path)
    sheet = workbook.active
    data = []

    for row in sheet.iter_rows(min_row=2, values_only=True):  # Skip header
        # Skip if row is empty or doesn't have at least 4 items
        if not row or len(row) < 4:
            continue

        # Unpack safely
        name, email, company, status = row[:4]

        # Skip if any of the required fields are missing/blank
        if not all([name, email, company, status]):
            continue

        data.append({
            "name": name,
            "email": email,
            "company": company,
            "status": status
        })

    return data
